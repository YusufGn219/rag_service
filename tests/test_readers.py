import pytest

from rag_service import readers


def _pdf_with_text(path, text):
    """A minimal one-page PDF whose page really contains `text` (ASCII only)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


def test_kind_of():
    assert readers.kind_of("a/B.PDF") == "pdf"
    assert readers.kind_of("x.htm") == "html"
    assert readers.kind_of("x.tsv") == "csv"
    assert readers.kind_of("x.md") is None
    assert readers.kind_of("x.doc") is None


def test_unsupported_extension_raises(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("x")
    with pytest.raises(ValueError):
        readers.extract(f)


def test_pdf_with_text(tmp_path):
    f = tmp_path / "Rapor.pdf"
    _pdf_with_text(f, "Merhaba dunya hizmet kapsami ve fiyat listesi burada")
    ex = readers.extract(f)
    assert ex.kind == "pdf" and ex.pages == 1
    assert not ex.hard
    assert ex.markdown.startswith("# Rapor\n")
    assert "## Sayfa 1" in ex.markdown and "Merhaba dunya" in ex.markdown


def test_pdf_without_text_is_hard(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    f = tmp_path / "Tarama.pdf"
    with open(f, "wb") as fh:
        w.write(fh)
    ex = readers.extract(f)
    assert ex.hard and "taranmış" in ex.reason or "metin" in ex.reason


def test_broken_pdf_raises(tmp_path):
    f = tmp_path / "Bozuk.pdf"
    f.write_bytes(b"not a pdf at all")
    with pytest.raises(readers.ReadError):
        readers.extract(f)


def test_docx_headings_lists_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Giriş", level=1)
    d.add_paragraph("Düz paragraf.")
    d.add_paragraph("birinci madde", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Ad", "Değer"
    t.cell(1, 0).text, t.cell(1, 1).text = "a|b", "5"
    f = tmp_path / "Belge.docx"
    d.save(str(f))
    ex = readers.extract(f)
    assert not ex.hard
    assert "## Giriş" in ex.markdown  # heading 1 sits under the "# Belge" title
    assert "Düz paragraf." in ex.markdown
    assert "- birinci madde" in ex.markdown
    assert "| Ad | Değer |" in ex.markdown and "a\\|b" in ex.markdown


def test_xlsx_sheets_become_tables(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Envanter"
    ws.append(["Cihaz", "Adet"])
    ws.append(["Yazıcı", 3.0])
    ws.append([None, None])  # empty row is skipped
    wb.create_sheet("Boş")  # empty sheet is skipped
    f = tmp_path / "Env.xlsx"
    wb.save(str(f))
    ex = readers.extract(f)
    assert "## Envanter" in ex.markdown and "## Boş" not in ex.markdown
    assert "| Cihaz | Adet |" in ex.markdown and "| Yazıcı | 3 |" in ex.markdown
    assert ex.pages == 2


def test_xlsx_row_limit(tmp_path, monkeypatch):
    import openpyxl

    monkeypatch.setattr(readers, "MAX_SHEET_ROWS", 3)
    wb = openpyxl.Workbook()
    for i in range(10):
        wb.active.append([f"r{i}"])
    f = tmp_path / "Uzun.xlsx"
    wb.save(str(f))
    ex = readers.extract(f)
    assert "r2" in ex.markdown and "r5" not in ex.markdown
    assert "kesildi" in ex.markdown


def test_pptx_slides_titles_and_notes(tmp_path):
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Planlama"
    s.placeholders[1].text = "ilk madde"
    s.notes_slide.notes_text_frame.text = "konuşmacı notu"
    f = tmp_path / "Sunum.pptx"
    p.save(str(f))
    ex = readers.extract(f)
    assert "## Slayt 1: Planlama" in ex.markdown
    assert "- ilk madde" in ex.markdown and "Not: konuşmacı notu" in ex.markdown
    assert ex.pages == 1


def test_csv_semicolon_and_turkish_encoding(tmp_path):
    f = tmp_path / "Liste.csv"
    f.write_bytes("ad;şehir\nAli;İzmir\n".encode("cp1254"))
    ex = readers.extract(f)
    assert "| ad | şehir |" in ex.markdown and "| Ali | İzmir |" in ex.markdown


def test_html_structure_and_noise(tmp_path):
    f = tmp_path / "Sayfa.html"
    f.write_text(
        "<html><head><title>Başlık</title><style>x{}</style></head><body>"
        "<nav>Menü</nav><h2>Bölüm</h2><p>Merhaba <b>dünya</b></p>"
        "<ul><li>bir</li><li>iki</li></ul>"
        "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "<script>alert(1)</script><!-- yorum --></body></html>",
        encoding="utf-8",
    )
    ex = readers.extract(f)
    md = ex.markdown
    assert "### Bölüm" in md
    assert "Merhaba dünya" in md
    assert "- bir" in md and "- iki" in md
    assert "| A | B |" in md and "| 1 | 2 |" in md
    for noise in ("Menü", "alert", "yorum", "x{}"):
        assert noise not in md


def test_txt_and_empty(tmp_path):
    f = tmp_path / "Not.txt"
    f.write_bytes("satır bir\r\nsatır iki".encode("utf-8"))
    assert "satır bir\nsatır iki" in readers.extract(f).markdown
    e = tmp_path / "Boş.txt"
    e.write_text("   \n")
    ex = readers.extract(e)
    assert ex.hard and ex.markdown == "# Boş\n"


def test_pdf_reflow_glues_one_word_per_line():
    text = "Cloud\n \nforensics\n \nhas\n \ngained\n \nimportance."
    assert readers._reflow(text) == "Cloud forensics has gained importance."


def test_pdf_reflow_keeps_normal_lines():
    text = "Bu satır yeterince uzun bir cümle.\n\nİkinci paragraf da uzun bir cümle."
    assert readers._reflow(text) == text
